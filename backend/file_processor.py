import os
import uuid
from pathlib import Path

UPLOAD_DIR = Path(__file__).resolve().parent / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)

TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".py", ".js", ".jsx", ".ts", ".tsx",
    ".json", ".csv", ".html", ".htm", ".css", ".scss", ".java", ".c",
    ".cpp", ".h", ".hpp", ".go", ".rs", ".rb", ".php", ".sql", ".yaml",
    ".yml", ".xml", ".sh", ".bat", ".ps1", ".env", ".ini", ".toml",
    ".log", ".rtf",
}

EXCEL_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
PPT_EXTENSIONS = {".pptx", ".pptm", ".potx"}


def save_upload(file_bytes: bytes, original_name: str) -> tuple[str, str]:
    ext = Path(original_name).suffix.lower()
    stored_name = f"{uuid.uuid4().hex}{ext}"
    dest = UPLOAD_DIR / stored_name
    dest.write_bytes(file_bytes)
    return str(dest), ext


def extract_text_from_file(file_path: str, original_name: str) -> str:
    path = Path(file_path)
    ext = Path(original_name).suffix.lower()

    if ext in TEXT_EXTENSIONS or ext == "":
        try:
            return path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            return path.read_text(encoding="latin-1")

    if ext in EXCEL_EXTENSIONS:
        return _extract_excel(path)

    if ext in PPT_EXTENSIONS:
        return _extract_ppt(path)

    if ext == ".docx":
        return _extract_docx(path)

    size = path.stat().st_size
    return (
        f"[Binary file: {original_name}]\n"
        f"Type: {ext or 'unknown'} | Size: {size:,} bytes\n"
        f"Stored for reference. Text extraction not supported for this format."
    )


def _extract_excel(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return f"[Excel file: {path.name}] Install openpyxl to extract content."

    wb = load_workbook(path, read_only=True, data_only=True)
    lines = [f"=== Excel: {path.name} ==="]
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        lines.append(f"\n--- Sheet: {sheet_name} ---")
        row_count = 0
        for row in ws.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                cells = [str(c) if c is not None else "" for c in row]
                lines.append("\t".join(cells))
                row_count += 1
                if row_count >= 100:
                    lines.append("... (truncated at 100 rows)")
                    break
    wb.close()
    return "\n".join(lines)


def _extract_ppt(path: Path) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        return f"[PowerPoint file: {path.name}] Install python-pptx to extract content."

    prs = Presentation(path)
    lines = [f"=== PowerPoint: {path.name} ==="]
    for i, slide in enumerate(prs.slides, 1):
        lines.append(f"\n--- Slide {i} ---")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                lines.append(shape.text.strip())
    return "\n".join(lines)


def _extract_docx(path: Path) -> str:
    try:
        from docx import Document as DocxDocument
    except ImportError:
        return f"[Word file: {path.name}] Install python-docx to extract content."

    doc = DocxDocument(path)
    lines = [f"=== Word: {path.name} ==="]
    for para in doc.paragraphs:
        if para.text.strip():
            lines.append(para.text)
    return "\n".join(lines)


def delete_file(file_path: str | None) -> None:
    if file_path and os.path.exists(file_path):
        os.remove(file_path)
